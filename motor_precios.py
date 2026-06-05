import pandas as pd
import numpy as np
from sqlalchemy import create_engine
import urllib.parse
import os
from datetime import datetime
from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

# 1. Cargar las variables del archivo .env
load_dotenv()

# 2. Configuración para LocalDB
DB_HOST = os.getenv("DB_HOST", "(localdb)\\MSSQLLocalDB")
DB_NAME = os.getenv("DB_NAME", "GestorPrecios")

params = urllib.parse.quote_plus(
    f"DRIVER={{ODBC Driver 17 for SQL Server}};"
    f"SERVER={DB_HOST};"
    f"DATABASE={DB_NAME};"
    f"Trusted_Connection=yes;"
)
engine = create_engine(f"mssql+pyodbc:///?odbc_connect={params}")

def exportar_excel_categorizado(df_flat: pd.DataFrame, filepath: str, titulo_general: str, dict_nombres: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Precios"

    # Paleta Corporativa
    font_titulo = Font(bold=True, size=14, color="FFFFFF")
    fill_titulo = PatternFill("solid", fgColor="0F172A")
    align_center = Alignment(horizontal="center", vertical="center")

    font_cat = Font(bold=True, size=12, color="FFFFFF")
    fill_cat = PatternFill("solid", fgColor="1E40AF")

    font_header_comercial = Font(bold=True, color="1E40AF")
    font_header_base = Font(bold=True, color="FFFFFF")
    fill_header_base = PatternFill("solid", fgColor="0F172A")

    # 1. Título Superior
    ws.append([titulo_general])
    ws.cell(row=ws.max_row, column=1).font = font_titulo
    ws.cell(row=ws.max_row, column=1).fill = fill_titulo
    ws.cell(row=ws.max_row, column=1).alignment = align_center

    # Requerimiento: Ordenamiento alfabético estricto tras la generación del Pivot
    if 'LINEA' in df_flat.columns:
        if 'Desc. item' in df_flat.columns:
            # Ordena primero por categoría y luego de la A a la Z por producto
            df_flat = df_flat.sort_values(by=['LINEA', 'Desc. item'], ascending=[True, True])
        else:
            df_flat = df_flat.sort_values('LINEA')
        lineas = df_flat['LINEA'].dropna().unique()
    else:
        lineas = ["GENERAL"]
        df_flat['LINEA'] = "GENERAL"

    for linea in lineas:
        df_linea = df_flat[df_flat['LINEA'] == linea].copy()
        if 'LINEA' in df_linea.columns:
            df_linea.drop(columns=['LINEA'], inplace=True)
        df_linea.rename(columns={'porcentaje_iva': 'IVA', 'U.M. Inv.': 'U.M.'}, inplace=True)

        headers = list(df_linea.columns)

        # Inyectar Separador de Categoría
        ws.append([str(linea).upper()])
        cat_row = ws.max_row
        ws.cell(row=cat_row, column=1).font = font_cat
        ws.cell(row=cat_row, column=1).fill = fill_cat
        ws.merge_cells(start_row=cat_row, start_column=1, end_row=cat_row, end_column=len(headers))

        # Dual Header: Fila 1 (Nombres Comerciales)
        super_headers = [dict_nombres.get(str(h), "") for h in headers]
        ws.append(super_headers)
        sup_row = ws.max_row
        for col_idx in range(1, len(super_headers) + 1):
            cell = ws.cell(row=sup_row, column=col_idx)
            cell.font = font_header_comercial
            cell.alignment = align_center

        # Dual Header: Fila 2 (Códigos y Columnas Base)
        ws.append(headers)
        head_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=head_row, column=col_idx)
            cell.font = font_header_base
            cell.fill = fill_header_base
            cell.alignment = align_center

        # Insertar Datos (Aplicando Formato Contable)
        for r in dataframe_to_rows(df_linea, index=False, header=False):
            ws.append(r)
            data_row = ws.max_row
            # Requerimiento 3: Formato Moneda dinámico
            for col_idx, val in enumerate(r, 1):
                cell = ws.cell(row=data_row, column=col_idx)
                if str(headers[col_idx-1]).isdigit() or headers[col_idx-1] == 'Precio con impuesto':
                    if isinstance(val, (int, float)):
                        cell.number_format = '"$"#,##0'

        ws.append([]) # Espacio entre categorías

    # Combinar el Título General a lo ancho de toda la tabla
    total_cols = len(df_flat.columns) if not df_flat.empty else 5
    # Título General está siempre en la fila 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)

    # Auto-ajuste de Columnas Blindado (ignora MergedCells)
    for col in ws.columns:
        max_length = 0
        column_letter = get_column_letter(col[0].column)
        for cell in col:
            if type(cell).__name__ == 'MergedCell':
                continue
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    # Ajuste de longitud si tiene formato de moneda
                    val_str = str(cell.value)
                    if cell.number_format == '"$"#,##0':
                        val_str = f"$ {val_str}"
                    max_length = max(max_length, len(val_str))
            except:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 45)

    wb.save(filepath)

def analizar_reportes(df_precios: pd.DataFrame, df_inventario: pd.DataFrame):
    """
    Retorna la lista de productos que fallan las reglas y NO son inmunes.
    """
    df_precios['Item'] = df_precios['Item'].astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
    df_inventario['Item'] = df_inventario['Item'].astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
    
    columnas_inventario = df_inventario.columns.difference(df_precios.columns).tolist()
    columnas_inventario.append('Item')
    df_inventario_limpio = df_inventario[columnas_inventario]
    
    df_merged = pd.merge(df_precios, df_inventario_limpio, on='Item', how='left')

    # Filtro Basura "NO-" (Requerimiento 1)
    if 'Desc. item' in df_merged.columns:
        no_mask = df_merged['Desc. item'].astype(str).str.strip().str.upper().str.startswith('NO-')
        df_merged = df_merged[~no_mask]

    if 'Lista' in df_merged.columns:
        df_merged['Lista'] = df_merged['Lista'].astype(str).str.strip()
        df_merged = df_merged[~df_merged['Lista'].isin({'01', '02', '1', '2', '13', '15', '101', '102'})]

    try:
        df_manuales = pd.read_sql("SELECT item FROM Items_Manuales", engine)
        items_inmunes = set(df_manuales['item'].astype(str).str.strip())
    except:
        items_inmunes = set()

    cant_disp = pd.to_numeric(df_merged.get('Cant. disponible', 0), errors='coerce').fillna(0)
    costo_prom = pd.to_numeric(df_merged.get('Costo prom. unit. (ins)', 0), errors='coerce').fillna(0)
    cond_a = (cant_disp > 0) & (costo_prom > 0)

    um_inv = df_merged.get('U.M. Inv.', pd.Series(dtype=str)).astype(str).str.strip().str.upper()
    peso = pd.to_numeric(df_merged.get('Peso en KILO', 0), errors='coerce').fillna(0)
    cond_b = ~((um_inv == 'KILO') & (peso != 1))

    # Fallan y no son inmunes
    mask_fail = ~(cond_a & cond_b)
    mask_not_immune = ~df_merged['Item'].isin(items_inmunes)
    
    df_omitidos = df_merged[mask_fail & mask_not_immune].copy()
    df_omitidos.drop_duplicates(subset=['Item'], inplace=True)
    
    # LA SOLUCIÓN: Limpiamos los NaN y los volvemos textos en blanco para JSON
    df_omitidos = df_omitidos.fillna("")
    
    omitidos_list = []
    for _, row in df_omitidos.iterrows():
        omitidos_list.append({
            'Item': row['Item'],
            'Desc. item': row.get('Desc. item', 'SIN DESCRIPCION')
        })
        
    return {
        "total_omitidos": len(omitidos_list),
        "omitidos": omitidos_list
    }

def procesar_reportes(df_precios: pd.DataFrame, df_inventario: pd.DataFrame, output_dir: str = "output", rescatados: list = None):
    os.makedirs(output_dir, exist_ok=True)

    # 1. Traer Diccionario de Listas
    dict_nombres_listas = {}
    try:
        df_listas = pd.read_sql("SELECT codigo_lista, nombre_lista FROM Maestro_Listas", engine)
        dict_nombres_listas = dict(zip(df_listas['codigo_lista'].astype(str).str.strip(), df_listas['nombre_lista']))
    except Exception as e:
        print(f"Alerta Maestro_Listas: {e}")

    # 2. Limpieza de .0 y Merge Principal
    df_precios['Item'] = df_precios['Item'].astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
    df_inventario['Item'] = df_inventario['Item'].astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
    
    columnas_inventario = df_inventario.columns.difference(df_precios.columns).tolist()
    columnas_inventario.append('Item')
    df_inventario_limpio = df_inventario[columnas_inventario]
    
    df_merged = pd.merge(df_precios, df_inventario_limpio, on='Item', how='left')

    # Filtro "NO-"
    if 'Desc. item' in df_merged.columns:
        no_mask = df_merged['Desc. item'].astype(str).str.strip().str.upper().str.startswith('NO-')
        df_merged = df_merged[~no_mask]

    # 3. Impuestos
    df_merged['Grupo impositivo'] = df_merged.get('Grupo impositivo', pd.Series("")).astype(str).str.strip()
    try:
        df_impuestos = pd.read_sql("SELECT codigo_impositivo, porcentaje_iva FROM Maestro_Impuestos", engine)
        df_impuestos['codigo_impositivo'] = df_impuestos['codigo_impositivo'].astype(str).str.strip()
        df_merged = pd.merge(df_merged, df_impuestos, left_on='Grupo impositivo', right_on='codigo_impositivo', how='left')
        df_merged['porcentaje_iva'] = df_merged['porcentaje_iva'].fillna(0)
    except:
        df_merged['porcentaje_iva'] = 0

    # 4. Filtro Basura Listas
    if 'Lista' in df_merged.columns:
        df_merged['Lista'] = df_merged['Lista'].astype(str).str.strip()
        df_merged = df_merged[~df_merged['Lista'].isin({'01', '02', '1', '2', '13', '15', '101', '102'})]

    # 5. Reglas de Negocio y Rescatados
    try:
        df_manuales = pd.read_sql("SELECT item FROM Items_Manuales", engine)
        items_inmunes = set(df_manuales['item'].astype(str).str.strip())
    except:
        items_inmunes = set()

    # Inyectar rescatados
    if rescatados:
        items_inmunes.update([str(r).strip() for r in rescatados])

    cant_disp = pd.to_numeric(df_merged.get('Cant. disponible', 0), errors='coerce').fillna(0)
    costo_prom = pd.to_numeric(df_merged.get('Costo prom. unit. (ins)', 0), errors='coerce').fillna(0)
    cond_a = (cant_disp > 0) & (costo_prom > 0)

    um_inv = df_merged.get('U.M. Inv.', pd.Series(dtype=str)).astype(str).str.strip().str.upper()
    peso = pd.to_numeric(df_merged.get('Peso en KILO', 0), errors='coerce').fillna(0)
    cond_b = ~((um_inv == 'KILO') & (peso != 1))

    df_merged = df_merged[df_merged['Item'].isin(items_inmunes) | (cond_a & cond_b)]

    # Requerimiento 2: Ordenamiento Alfabético
    if 'LINEA' in df_merged.columns and 'Desc. item' in df_merged.columns:
        df_merged.sort_values(by=['LINEA', 'Desc. item'], ascending=[True, True], inplace=True)

    # 6. Pivot Table
    index_cols = ['LINEA', 'Item', 'Desc. item', 'U.M. Inv.', 'Peso en KILO', 'porcentaje_iva']
    for col in index_cols:
        df_merged[col] = df_merged.get(col, pd.Series("SIN DATO")).fillna("SIN DATO")

    df_merged['Precio con impuesto'] = df_merged.get('Precio con impuesto', 0)

    pivot = df_merged.pivot_table(
        index=index_cols,
        columns='Lista',
        values='Precio con impuesto',
        aggfunc='first'
    )

    # Requerimiento 4: Orden Estricto de Columnas
    orden_oficial = ['11', '12', '03', '04', '05', '06', '07', '08', '09', '10', '14', '16']
    # Mantener solo las que realmente vinieron en el excel para no crear columnas vacías
    columnas_presentes = [col for col in orden_oficial if col in pivot.columns]
    
    if columnas_presentes:
        pivot = pivot[columnas_presentes]
    else:
        # Si por casualidad suben un excel sin ninguna de esas listas, pasamos de largo
        pass 

    archivos_generados = []
    fecha_actual = datetime.now()
    fecha_header = fecha_actual.strftime('%d MAYO %Y').upper()

    # --- REPORTE GENERAL ---
    if not pivot.empty:
        filepath_general = os.path.join(output_dir, "reporte_GENERAL.xlsx")
        titulo_general = f"LISTA DE PRECIOS GENERAL - {fecha_header}"
        exportar_excel_categorizado(pivot.reset_index(), filepath_general, titulo_general, dict_nombres_listas)
        archivos_generados.append(filepath_general)

    # --- REPORTES ASESORES ---
    try:
        df_asesores = pd.read_sql("SELECT nombre_asesor, listas_asignadas FROM Maestro_Asesores", engine)
    except:
        df_asesores = pd.DataFrame()

    historico_records = []

    for _, asesor in df_asesores.iterrows():
        nombre = str(asesor['nombre_asesor']).strip()
        listas_asignadas = [l.strip() for l in str(asesor['listas_asignadas']).split(',') if l.strip()]

        # Mantener el orden oficial también para el asesor
        columnas_disponibles = [col for col in pivot.columns if str(col) in listas_asignadas]
        
        if not columnas_disponibles:
            continue

        pivot_asesor = pivot[columnas_disponibles].copy()
        pivot_asesor.dropna(how='all', inplace=True)

        if pivot_asesor.empty:
            continue

        filepath = os.path.join(output_dir, f"reporte_{nombre}.xlsx")
        titulo_asesor = f"LISTA DE PRECIOS - {nombre} - {fecha_header}"
        
        exportar_excel_categorizado(pivot_asesor.reset_index(), filepath, titulo_asesor, dict_nombres_listas)
        archivos_generados.append(filepath)

        # Auditoría
        melted = pivot_asesor.reset_index().melt(
            id_vars=['Item', 'Desc. item'],
            value_vars=columnas_disponibles,
            var_name='lista',
            value_name='precio_con_impuesto'
        ).dropna(subset=['precio_con_impuesto'])

        for _, m_row in melted.iterrows():
            historico_records.append({
                'fecha': fecha_actual,
                'item': m_row['Item'],
                'descripcion': m_row['Desc. item'],
                'precio_con_impuesto': m_row['precio_con_impuesto'],
                'lista': m_row['lista'],
                'asesor_destino': nombre
            })

    if historico_records:
        try:
            pd.DataFrame(historico_records).to_sql('Historico_Generaciones', engine, if_exists='append', index=False)
        except:
            pass

    return archivos_generados